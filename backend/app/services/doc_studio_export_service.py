"""Document Studio export (PDF / DOCX / HTML / Markdown) and import (DOCX / PDF / MD / TXT).

The markdown dialect is what ``tiptap-markdown`` emits: ATX headings, ``-``/``1.``
lists, ``- [ ]`` tasks, ``>`` quotes, fenced code, pipe tables, ``![alt](src)``
images, ``**bold**`` / ``*em*`` / `` `code` `` inline marks and ``---`` rules.
"""

from __future__ import annotations

import base64
import html
import io
import logging
import re
from pathlib import Path
from dataclasses import dataclass, field

from app.models.doc_document import DocAsset, DocDocument
from app.services.doc_studio_service import ASSET_URL_PREFIX, DocStudioService
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)


BRAND_NAVY = "#07111f"
BRAND_SKY = "#38bdf8"
BRAND_SLATE = "#475569"
BRAND_LINE = "WATER WORKFORCE 360 · NYSAWWA · ONE WATER WORKFORCE"

# Logo for printables — prefer the light-background lockup shipped with the backend image.
_BRAND_LOGO_CANDIDATES = (
    Path(__file__).resolve().parent.parent / "assets" / "brand" / "workforce-360-logo.png",
    Path(__file__).resolve().parents[3] / "frontend" / "public" / "workforce-360-logo.png",
)


def brand_logo_path() -> Path | None:
    for candidate in _BRAND_LOGO_CANDIDATES:
        try:
            if candidate.is_file():
                return candidate
        except OSError:
            continue
    return None


def brand_logo_data_uri() -> str | None:
    path = brand_logo_path()
    if not path:
        return None
    try:
        data = path.read_bytes()
    except OSError:
        return None
    return "data:image/png;base64," + base64.b64encode(data).decode("ascii")



# ── Markdown block parser (small, dependency-free) ──────────────────────────


@dataclass
class Block:
    kind: (
        str  # heading | paragraph | bullets | numbers | tasks | quote | code | table | image | rule | pagebreak
    )
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    checked: list[bool] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    src: str = ""
    alt: str = ""
    width_pct: int | None = None


_IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)\s*$")
_INLINE_IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)")
_TASK_RE = re.compile(r"^\s*[-*+]\s+\[( |x|X)\]\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUM_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_IMG_WIDTH_RE = re.compile(r"^width=(\d{1,3})$", re.I)
_IMG_FLOAT_RE = re.compile(r"^float=(left|right|none)$", re.I)


def split_image_alt(raw_alt: str) -> tuple[str, int | None]:
    """Parse `caption|width=60|float=left` → (caption, width_pct)."""
    if not raw_alt:
        return "", None
    caption_parts: list[str] = []
    width: int | None = None
    for part in (p.strip() for p in raw_alt.split("|")):
        wm = _IMG_WIDTH_RE.match(part)
        if wm:
            parsed = int(wm.group(1))
            if 5 <= parsed <= 100:
                width = parsed
            continue
        if _IMG_FLOAT_RE.match(part):
            continue
        caption_parts.append(part)
    return "|".join(caption_parts), width


def expand_text_with_images(text: str) -> list[Block]:
    """Split a text run that may contain inline `![alt](src)` into paragraph/image blocks."""
    text = (text or "").strip()
    if not text:
        return []
    if not _INLINE_IMG_RE.search(text):
        return [Block(kind="paragraph", text=text)]
    out: list[Block] = []
    pos = 0
    for m in _INLINE_IMG_RE.finditer(text):
        before = text[pos : m.start()].strip()
        if before:
            out.append(Block(kind="paragraph", text=before))
        caption, width_pct = split_image_alt(m.group(1))
        out.append(Block(kind="image", alt=caption, src=m.group(2), width_pct=width_pct))
        pos = m.end()
    after = text[pos:].strip()
    if after:
        out.append(Block(kind="paragraph", text=after))
    return out


def normalize_markdown(markdown: str) -> str:
    """Normalize TipTap hard-breaks so images/headings parse cleanly.

    TipTap often emits a single ``\\`` before a newline (hard break), and sometimes
    glues ``\\## Heading`` onto the same line as an image. Turn those into real
    newlines so block parsing and export stay consistent.

    Important: do **not** rewrite a single ``\\#`` — TipTap escapes ``#`` that way
    inside table cells (e.g. a ``#`` column). Treating it as a heading hard-break
    destroys pipe tables in export.
    """
    text = (markdown or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\\\n", "\n", text)
    # Only ATX h2–h6 glued after a hard break (``\## Title``), never ``\#``.
    text = re.sub(r"\\(#{2,6}\s)", r"\n\1", text)
    return text


def unescape_md_cell(text: str) -> str:
    """Undo TipTap markdown escapes commonly found in table cells / inline runs."""
    if not text:
        return ""
    return (
        text.replace("\\#", "#")
        .replace("\\|", "|")
        .replace("\\>", ">")
        .replace("\\-", "-")
        .replace("\\*", "*")
        .replace("\\\\", "\\")
    )


def parse_markdown(markdown: str) -> list[Block]:
    lines = normalize_markdown(markdown).split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)
    para: list[str] = []

    def flush_para() -> None:
        nonlocal para
        if para:
            text = " ".join(s.strip() for s in para)
            blocks.extend(expand_text_with_images(text))
            para = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_para()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_para()
            i += 1
            code: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block(kind="code", text="\n".join(code)))
            continue

        m = _HEADING_RE.match(stripped)
        if m:
            flush_para()
            blocks.append(Block(kind="heading", level=len(m.group(1)), text=m.group(2)))
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush_para()
            blocks.append(Block(kind="rule"))
            i += 1
            continue

        if stripped in ("--- Page Break ---", "<!-- pagebreak -->") or re.match(
            r"^<!--\s*pagebreak\s*-->$", stripped, re.I
        ):
            flush_para()
            blocks.append(Block(kind="pagebreak"))
            i += 1
            continue

        m = _IMG_RE.match(stripped)
        if m:
            flush_para()
            caption, width_pct = split_image_alt(m.group(1))
            blocks.append(Block(kind="image", alt=caption, src=m.group(2), width_pct=width_pct))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1]):
            flush_para()
            rows: list[list[str]] = []

            def _split_row(row_line: str) -> list[str]:
                return [unescape_md_cell(c.strip()) for c in row_line.strip().strip("|").split("|")]

            rows.append(_split_row(stripped))
            i += 2  # skip header + separator
            while i < n and lines[i].strip().startswith("|"):
                if _TABLE_SEP_RE.match(lines[i]):
                    i += 1
                    continue
                rows.append(_split_row(lines[i]))
                i += 1
            blocks.append(Block(kind="table", rows=rows))
            continue

        if stripped.startswith(">"):
            flush_para()
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(Block(kind="quote", text=" ".join(q for q in quote if q)))
            continue

        if _TASK_RE.match(line):
            flush_para()
            items: list[str] = []
            checked: list[bool] = []
            while i < n and _TASK_RE.match(lines[i]):
                tm = _TASK_RE.match(lines[i])
                assert tm
                checked.append(tm.group(1).lower() == "x")
                items.append(tm.group(2))
                i += 1
            blocks.append(Block(kind="tasks", items=items, checked=checked))
            continue

        if _BULLET_RE.match(line):
            flush_para()
            items = []
            while i < n and _BULLET_RE.match(lines[i]) and not _TASK_RE.match(lines[i]):
                bm = _BULLET_RE.match(lines[i])
                assert bm
                items.append(bm.group(1))
                i += 1
            blocks.append(Block(kind="bullets", items=items))
            continue

        if _NUM_RE.match(line):
            flush_para()
            items = []
            while i < n and _NUM_RE.match(lines[i]):
                nm = _NUM_RE.match(lines[i])
                assert nm
                items.append(nm.group(1))
                i += 1
            blocks.append(Block(kind="numbers", items=items))
            continue

        para.append(line)
        i += 1

    flush_para()
    return blocks


# ── Inline formatting ───────────────────────────────────────────────────────

# Negative lookbehind so `![alt](src)` image markdown is not treated as a link.
_INLINE_LINK = re.compile(r"(?<!!)\[([^\]]+)\]\(([^)\s]+)\)")
_INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")
_INLINE_EM = re.compile(r"(?<![*\w])\*(?!\*)(.+?)(?<!\*)\*(?![*\w])|(?<!\w)_(.+?)_(?!\w)")
_INLINE_CODE = re.compile(r"`([^`]+)`")
_INLINE_STRIKE = re.compile(r"~~(.+?)~~")


def inline_to_html(text: str) -> str:
    """Markdown inline marks → minimal HTML (safe for ReportLab Paragraph and HTML export)."""
    # Images should already be split into blocks; strip any leftovers so link parsing
    # cannot turn `![alt](src)` into `!<a>alt</a>`.
    cleaned = _INLINE_IMG_RE.sub("", text or "")
    out = html.escape(cleaned, quote=False)
    out = _INLINE_CODE.sub(lambda m: f"<font face='Courier'>{m.group(1)}</font>", out)
    out = _INLINE_BOLD.sub(lambda m: f"<b>{m.group(1) or m.group(2)}</b>", out)
    out = _INLINE_EM.sub(lambda m: f"<i>{m.group(1) or m.group(2)}</i>", out)
    out = _INLINE_STRIKE.sub(lambda m: f"<strike>{m.group(1)}</strike>", out)
    out = _INLINE_LINK.sub(
        lambda m: f"<a href='{m.group(2)}' color='#0369a1'>{m.group(1)}</a>", out
    )
    return out


def inline_to_plain(text: str) -> str:
    cleaned = _INLINE_IMG_RE.sub("", text or "")
    out = _INLINE_LINK.sub(lambda m: m.group(1), cleaned)
    for rx in (_INLINE_BOLD, _INLINE_EM, _INLINE_CODE, _INLINE_STRIKE):
        out = rx.sub(lambda m: next(g for g in m.groups() if g is not None), out)
    return out


# ── Export service ──────────────────────────────────────────────────────────


class DocStudioExportService:
    def __init__(self, db: Session):
        self.db = db
        self.studio = DocStudioService(db)

    def _doc(self, scope: str, document_id: str) -> DocDocument:
        return self.studio._get_document_row(scope, document_id)

    def _asset_row(self, src: str) -> DocAsset | None:
        m = re.search(
            rf"{re.escape(ASSET_URL_PREFIX)}/([0-9a-fA-F-]{{36}})/file",
            src or "",
        )
        if not m:
            return None
        return self.db.query(DocAsset).filter(DocAsset.id == m.group(1)).first()

    def _asset_bytes(self, src: str) -> bytes | None:
        a = self._asset_row(src)
        return bytes(a.data) if a and a.data is not None else None

    def _asset_mime(self, src: str) -> str:
        a = self._asset_row(src)
        ctype = (a.content_type if a else None) or "image/png"
        if not ctype.startswith("image/"):
            return "image/png"
        return ctype

    def _image_data_uri(self, src: str) -> str | None:
        raw = self._asset_bytes(src)
        if not raw:
            return None
        mime = self._asset_mime(src)
        return f"data:{mime};base64," + base64.b64encode(raw).decode("ascii")

    def _reportlab_image(self, raw: bytes, src: str):
        """Build a ReportLab Image with a named buffer so format sniffing works."""
        from reportlab.platypus import Image

        mime = self._asset_mime(src)
        ext = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/jpg": ".jpg",
            "image/gif": ".gif",
            "image/webp": ".webp",
        }.get(mime, ".png")
        buf = io.BytesIO(raw)
        buf.name = f"asset{ext}"
        return Image(buf)

    # Markdown
    def export_markdown(self, scope: str, document_id: str) -> str:
        doc = self._doc(scope, document_id)
        body = doc.content_markdown or ""
        if body.lstrip().startswith("# "):
            return body
        return f"# {doc.title}\n\n{body}"

    # HTML (used for print preview and as the DOCX/PDF fallback)
    def export_html(self, scope: str, document_id: str) -> str:
        doc = self._doc(scope, document_id)
        parts: list[str] = [
            "<!doctype html><html><head><meta charset='utf-8'>",
            f"<title>{html.escape(doc.title or 'Document')}</title>",
            "<style>body{font-family:Inter,system-ui,sans-serif;color:#0f172a;max-width:52rem;margin:2rem auto;padding:0 1rem;line-height:1.55}"
            f"h1{{color:{BRAND_NAVY}}} h2{{border-bottom:2px solid {BRAND_SKY};padding-bottom:.25rem}} "
            "table{border-collapse:collapse;width:100%} td,th{border:1px solid #cbd5e1;padding:.35rem .5rem} th{background:#e0f2fe;text-align:left} "
            "blockquote{border-left:4px solid #38bdf8;margin:0;padding:.25rem 1rem;color:#334155;background:#f0f9ff} "
            "pre{background:#0f172a;color:#e2e8f0;padding:.75rem;border-radius:.5rem;overflow:auto} img{max-width:100%} "
            ".brand{font-size:.8rem;letter-spacing:.12em;text-transform:uppercase;color:#0369a1;font-weight:600;margin:.5rem 0 0}"
            ".brand-lockup{display:flex;flex-direction:column;align-items:flex-start;gap:.15rem;margin:0 0 1.35rem}"
            ".brand-lockup img{height:96px;width:auto;max-width:min(100%,360px);display:block}"
            "@media print{.brand-lockup img{height:88px}}</style></head><body>",
        ]
        logo_uri = brand_logo_data_uri()
        if logo_uri:
            parts.append(
                "<div class='brand-lockup'>"
                f"<img src='{logo_uri}' alt='Water Workforce 360'/>"
                f"<p class='brand'>{html.escape(BRAND_LINE)}</p>"
                "</div>"
            )
        else:
            parts.append(f"<p class='brand'>{html.escape(BRAND_LINE)}</p>")
        parts.append(f"<h1>{html.escape(doc.title or '')}</h1>")
        for b in parse_markdown(doc.content_markdown or ""):
            if b.kind == "heading":
                lvl = min(6, max(1, b.level + 1))
                parts.append(f"<h{lvl}>{inline_to_html(b.text)}</h{lvl}>")
            elif b.kind == "paragraph":
                parts.append(f"<p>{inline_to_html(b.text)}</p>")
            elif b.kind == "bullets":
                parts.append(
                    "<ul>" + "".join(f"<li>{inline_to_html(t)}</li>" for t in b.items) + "</ul>"
                )
            elif b.kind == "numbers":
                parts.append(
                    "<ol>" + "".join(f"<li>{inline_to_html(t)}</li>" for t in b.items) + "</ol>"
                )
            elif b.kind == "tasks":
                lis = "".join(
                    f"<li>{'☑' if c else '☐'} {inline_to_html(t)}</li>"
                    for t, c in zip(b.items, b.checked, strict=False)
                )
                parts.append(f"<ul style='list-style:none;padding-left:0'>{lis}</ul>")
            elif b.kind == "quote":
                parts.append(f"<blockquote>{inline_to_html(b.text)}</blockquote>")
            elif b.kind == "code":
                parts.append(f"<pre><code>{html.escape(b.text)}</code></pre>")
            elif b.kind == "rule":
                parts.append("<hr/>")
            elif b.kind == "pagebreak":
                parts.append("<div style='page-break-before:always;break-before:page'></div>")
            elif b.kind == "image":
                style = ""
                if b.width_pct:
                    style = f" style='width:{b.width_pct}%;height:auto'"
                # Embed bytes so Print view / downloaded HTML work without a live
                # authenticated session for `/assets/.../file`.
                uri = self._image_data_uri(b.src) or b.src
                parts.append(
                    f"<img src='{html.escape(uri, quote=True)}' alt='{html.escape(b.alt)}'{style}/>"
                )
            elif b.kind == "table" and b.rows:
                head = "".join(f"<th>{inline_to_html(c)}</th>" for c in b.rows[0])
                body_rows = "".join(
                    "<tr>" + "".join(f"<td>{inline_to_html(c)}</td>" for c in r) + "</tr>"
                    for r in b.rows[1:]
                )
                parts.append(
                    f"<table><thead><tr>{head}</tr></thead><tbody>{body_rows}</tbody></table>"
                )
        parts.append("</body></html>")
        return "".join(parts)

    # PDF (ReportLab)
    def export_pdf(self, scope: str, document_id: str) -> bytes:
        doc = self._doc(scope, document_id)
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            HRFlowable,
            Image,
            ListFlowable,
            ListItem,
            PageBreak,  # noqa: F401
            Paragraph,
            Preformatted,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        styles = getSampleStyleSheet()
        base = ParagraphStyle(
            "ww-body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15
        )
        h_styles = {
            1: ParagraphStyle(
                "ww-h1",
                parent=styles["Heading1"],
                textColor=colors.HexColor(BRAND_NAVY),
                spaceBefore=10,
                spaceAfter=6,
            ),
            2: ParagraphStyle(
                "ww-h2",
                parent=styles["Heading2"],
                textColor=colors.HexColor(BRAND_NAVY),
                spaceBefore=10,
                spaceAfter=4,
            ),
            3: ParagraphStyle(
                "ww-h3",
                parent=styles["Heading3"],
                textColor=colors.HexColor(BRAND_SLATE),
                spaceBefore=8,
                spaceAfter=3,
            ),
        }
        quote_style = ParagraphStyle(
            "ww-quote",
            parent=base,
            leftIndent=14,
            textColor=colors.HexColor("#334155"),
            borderPadding=(4, 6, 4, 6),
            backColor=colors.HexColor("#f0f9ff"),
            borderColor=colors.HexColor(BRAND_SKY),
            borderWidth=0,
        )
        code_style = ParagraphStyle(
            "ww-code",
            parent=styles["Code"],
            fontSize=8.5,
            leading=11,
            backColor=colors.HexColor("#f1f5f9"),
            borderPadding=(4, 6, 4, 6),
        )
        brand_style = ParagraphStyle(
            "ww-brand",
            parent=base,
            fontSize=9,
            textColor=colors.HexColor("#0369a1"),
            alignment=TA_LEFT,
            spaceAfter=2,
        )
        title_style = ParagraphStyle(
            "ww-title",
            parent=styles["Title"],
            textColor=colors.HexColor(BRAND_NAVY),
            alignment=TA_LEFT,
            fontSize=22,
            leading=26,
            spaceAfter=4,
        )

        buf = io.BytesIO()
        pdf = SimpleDocTemplate(
            buf,
            pagesize=letter,
            leftMargin=0.9 * inch,
            rightMargin=0.9 * inch,
            topMargin=0.8 * inch,
            bottomMargin=0.8 * inch,
            title=doc.title or "Document",
            author="Water Workforce 360",
        )
        story: list = []
        logo_path = brand_logo_path()
        if logo_path:
            try:
                # ~3.25" wide lockup; height from intrinsic 1200×867 aspect ratio.
                logo_w = 3.25 * inch
                logo_h = logo_w * (867 / 1200)
                story.append(Image(str(logo_path), width=logo_w, height=logo_h))
                story.append(Spacer(1, 6))
            except Exception:
                logger.exception("Failed to embed brand logo in PDF export")
        story.extend(
            [
                Paragraph(BRAND_LINE, brand_style),
                Paragraph(html.escape(doc.title or "Document"), title_style),
                HRFlowable(
                    width="100%", thickness=2, color=colors.HexColor(BRAND_SKY), spaceAfter=10
                ),
            ]
        )
        avail_w = letter[0] - pdf.leftMargin - pdf.rightMargin

        for b in parse_markdown(doc.content_markdown or ""):
            if b.kind == "heading":
                story.append(
                    Paragraph(inline_to_html(b.text), h_styles.get(min(b.level, 3), h_styles[3]))
                )
            elif b.kind == "paragraph":
                story.append(Paragraph(inline_to_html(b.text), base))
                story.append(Spacer(1, 4))
            elif b.kind in ("bullets", "numbers", "tasks"):
                if b.kind == "tasks":
                    items = [
                        ListItem(
                            Paragraph(("☑ " if c else "☐ ") + inline_to_html(t), base), leftIndent=6
                        )
                        for t, c in zip(b.items, b.checked, strict=False)
                    ]
                    story.append(ListFlowable(items, bulletType="bullet", start="", leftIndent=10))
                else:
                    items = [ListItem(Paragraph(inline_to_html(t), base)) for t in b.items]
                    story.append(
                        ListFlowable(
                            items,
                            bulletType="1" if b.kind == "numbers" else "bullet",
                            leftIndent=14,
                        )
                    )
                story.append(Spacer(1, 4))
            elif b.kind == "quote":
                story.append(Paragraph(inline_to_html(b.text), quote_style))
                story.append(Spacer(1, 6))
            elif b.kind == "code":
                story.append(Preformatted(b.text, code_style))
                story.append(Spacer(1, 6))
            elif b.kind == "rule":
                story.append(
                    HRFlowable(
                        width="100%",
                        thickness=0.6,
                        color=colors.HexColor("#cbd5e1"),
                        spaceBefore=6,
                        spaceAfter=6,
                    )
                )
            elif b.kind == "pagebreak":
                story.append(PageBreak())
            elif b.kind == "image":
                raw = self._asset_bytes(b.src)
                if raw:
                    try:
                        img = self._reportlab_image(raw, b.src)
                        ratio = img.imageHeight / float(img.imageWidth or 1)
                        pct = (b.width_pct or 100) / 100.0
                        w = avail_w * pct
                        # Intrinsic size is often pixels-as-points; don't upscale tiny assets.
                        if img.imageWidth and img.imageWidth < w:
                            w = float(img.imageWidth)
                        img.drawWidth = w
                        img.drawHeight = w * ratio
                        story.append(img)
                        story.append(Spacer(1, 6))
                    except Exception as exc:  # pragma: no cover
                        logger.warning("PDF image skipped (%s): %s", b.src, exc)
                        if b.alt:
                            story.append(
                                Paragraph(f"<i>[Image: {html.escape(b.alt)}]</i>", base)
                            )
                elif b.alt:
                    story.append(Paragraph(f"<i>[Image: {html.escape(b.alt)}]</i>", base))
                else:
                    story.append(Paragraph("<i>[Image missing]</i>", base))
            elif b.kind == "table" and b.rows:
                ncols = max(len(r) for r in b.rows)
                data = [
                    [Paragraph(inline_to_html(c), base) for c in (r + [""] * (ncols - len(r)))]
                    for r in b.rows
                ]
                t = Table(data, colWidths=[avail_w / ncols] * ncols, repeatRows=1)
                t.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0f2fe")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 5),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ]
                    )
                )
                story.append(t)
                story.append(Spacer(1, 8))

        def _footer(canvas, d):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#64748b"))
            canvas.drawString(
                d.leftMargin, 0.5 * inch, f"{doc.title or 'Document'} · v{doc.version_no}"
            )
            canvas.drawRightString(letter[0] - d.rightMargin, 0.5 * inch, f"Page {d.page}")
            canvas.restoreState()

        pdf.build(story, onFirstPage=_footer, onLaterPages=_footer)
        return buf.getvalue()

    # DOCX (python-docx)
    def export_docx(self, scope: str, document_id: str) -> bytes:
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required for DOCX export") from exc

        doc = self._doc(scope, document_id)
        d = Document()
        logo_path = brand_logo_path()
        if logo_path:
            try:
                d.add_picture(str(logo_path), width=Inches(3.25))
            except Exception:
                logger.exception("Failed to embed brand logo in DOCX export")
        brand = d.add_paragraph()
        run = brand.add_run(BRAND_LINE)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
        d.add_heading(doc.title or "Document", level=0)

        def add_inline(paragraph, text: str) -> None:
            # Very small inline-mark renderer: **bold**, *em*, `code`.
            pos = 0
            token_re = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![*\w])\*(?!\*).+?\*(?![*\w]))")
            for m in token_re.finditer(text):
                if m.start() > pos:
                    paragraph.add_run(inline_to_plain(text[pos : m.start()]))
                tok = m.group(0)
                if tok.startswith("**"):
                    r = paragraph.add_run(tok[2:-2])
                    r.bold = True
                elif tok.startswith("`"):
                    r = paragraph.add_run(tok[1:-1])
                    r.font.name = "Courier New"
                else:
                    r = paragraph.add_run(tok[1:-1])
                    r.italic = True
                pos = m.end()
            if pos < len(text):
                paragraph.add_run(inline_to_plain(text[pos:]))

        for b in parse_markdown(doc.content_markdown or ""):
            if b.kind == "heading":
                d.add_heading(inline_to_plain(b.text), level=min(b.level, 4))
            elif b.kind == "paragraph":
                add_inline(d.add_paragraph(), b.text)
            elif b.kind == "bullets":
                for t in b.items:
                    add_inline(d.add_paragraph(style="List Bullet"), t)
            elif b.kind == "numbers":
                for t in b.items:
                    add_inline(d.add_paragraph(style="List Number"), t)
            elif b.kind == "tasks":
                for t, c in zip(b.items, b.checked, strict=False):
                    add_inline(d.add_paragraph(), ("☑ " if c else "☐ ") + t)
            elif b.kind == "quote":
                p = d.add_paragraph(style="Intense Quote")
                add_inline(p, b.text)
            elif b.kind == "code":
                p = d.add_paragraph()
                r = p.add_run(b.text)
                r.font.name = "Courier New"
                r.font.size = Pt(9)
            elif b.kind == "rule":
                d.add_paragraph("―" * 30)
            elif b.kind == "pagebreak":
                d.add_page_break()
            elif b.kind == "image":
                raw = self._asset_bytes(b.src)
                if raw:
                    try:
                        pct = (b.width_pct or 100) / 100.0
                        mime = self._asset_mime(b.src)
                        ext = {
                            "image/png": ".png",
                            "image/jpeg": ".jpg",
                            "image/jpg": ".jpg",
                            "image/gif": ".gif",
                        }.get(mime, ".png")
                        buf = io.BytesIO(raw)
                        buf.name = f"asset{ext}"
                        d.add_picture(buf, width=Inches(6 * pct))
                    except Exception as exc:  # pragma: no cover
                        logger.warning("DOCX image skipped (%s): %s", b.src, exc)
                elif b.alt:
                    d.add_paragraph(f"[Image: {b.alt}]")
                else:
                    d.add_paragraph("[Image missing]")
            elif b.kind == "table" and b.rows:
                ncols = max(len(r) for r in b.rows)
                t = d.add_table(rows=len(b.rows), cols=ncols)
                t.style = "Light Grid Accent 1"
                for ri, r in enumerate(b.rows):
                    for ci in range(ncols):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        add_inline(cell.paragraphs[0], r[ci] if ci < len(r) else "")
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()


# ── Import: uploaded file → markdown ────────────────────────────────────────


def _docx_to_markdown(data: bytes) -> str:
    from docx import Document

    d = Document(io.BytesIO(data))
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        style = (p.style.name or "").lower() if p.style is not None else ""
        if not text:
            out.append("")
            continue
        if style.startswith("heading"):
            try:
                lvl = int(style.split()[-1])
            except ValueError:
                lvl = 2
            out.append("#" * max(1, min(lvl, 6)) + " " + text)
        elif style.startswith("title"):
            out.append("# " + text)
        elif "bullet" in style:
            out.append("- " + text)
        elif "number" in style:
            out.append("1. " + text)
        else:
            out.append(text)
    for t in d.tables:
        out.append("")
        rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
        if rows:
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("| " + " | ".join("---" for _ in rows[0]) + " |")
            for r in rows[1:]:
                out.append("| " + " | ".join(r) + " |")
    return "\n".join(out).strip() + "\n"


def _pdf_to_markdown(data: bytes) -> str:
    try:
        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                if txt.strip():
                    parts.append(txt.strip())
        return "\n\n".join(parts) + "\n"
    except Exception:  # pragma: no cover
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "").strip() for p in reader.pages) + "\n"


def file_to_markdown(filename: str, content_type: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx") or "wordprocessingml" in (content_type or ""):
        return _docx_to_markdown(data)
    if name.endswith(".pdf") or content_type == "application/pdf":
        return _pdf_to_markdown(data)
    if name.endswith((".md", ".markdown", ".txt")) or (content_type or "").startswith("text/"):
        return data.decode("utf-8", errors="replace")
    raise ValueError("Unsupported file type — upload .docx, .pdf, .md or .txt")


_BRAND_LINE = "WATER WORKFORCE 360"


def title_from_markdown(markdown: str, fallback: str) -> str:
    """Prefer the first H1 in the opening lines; otherwise the first real line."""
    lines = [ln.strip() for ln in (markdown or "").splitlines()]
    head = [ln for ln in lines[:20] if ln]
    for s in head:
        if s.startswith("# "):
            return s[2:].strip()[:200]
    for s in head:
        if s.upper().startswith(_BRAND_LINE):
            continue
        return s.lstrip("#").strip()[:120]
    return fallback
